#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Word 导出模块 - Markdown 简历 → DOCX
使用 python-docx，支持标题、正文、列表、粗体等样式
"""

import os
import sys
import re
import argparse
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

# 默认头像路径
SKILL_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_AVATAR = os.path.join(SKILL_DIR, "assets", "avatar.jpg")


def set_font(run, name="微软雅黑", size=10.5, bold=False, color=None):
    """设置字体"""
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    # 设置中文字体
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def add_inline_formatting(paragraph, text):
    """处理行内格式（**粗体**）"""
    parts = re.split(r'(\*\*[^*]+\*\*)', text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_font(run, bold=True)
        else:
            run = paragraph.add_run(part)
            set_font(run)


def md_to_docx(md_path: str, docx_path: str, avatar_path: str = None) -> bool:
    """Markdown → DOCX（带头像）"""
    if not os.path.exists(md_path):
        print(f"❌ 文件不存在: {md_path}", file=sys.stderr)
        return False

    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    doc = Document()

    # 页面设置
    section = doc.sections[0]
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    # 默认样式
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(10.5)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    # ===== 处理头部：姓名 + 联系方式 + 头像 =====
    name = ""
    contact = ""
    start_idx = 0

    # 提取 h1（姓名）
    for i, line in enumerate(lines):
        line = line.rstrip("\n")
        if line.startswith("# "):
            name = line[2:].strip()
            start_idx = i + 1
            break

    # 提取 h1 后面的第一个 p（联系方式）
    for i in range(start_idx, len(lines)):
        line = lines[i].rstrip("\n").strip()
        if line and not line.startswith("#") and line != "---":
            contact = line
            start_idx = i + 1
            break

    # 创建头部表格（1行2列）
    if name:
        table = doc.add_table(rows=1, cols=2)
        table.autofit = False
        # 设置列宽
        table.columns[0].width = Cm(13)
        table.columns[1].width = Cm(3.5)

        # 左侧：姓名 + 联系方式
        left_cell = table.cell(0, 0)
        left_cell.vertical_alignment = 1  # 居中
        p_name = left_cell.paragraphs[0]
        p_name.paragraph_format.space_after = Pt(4)
        run_name = p_name.add_run(name)
        set_font(run_name, size=20, bold=True)

        if contact:
            p_contact = left_cell.add_paragraph()
            p_contact.paragraph_format.space_after = Pt(0)
            run_contact = p_contact.add_run(contact)
            set_font(run_contact, size=9, color=(100, 100, 100))

        # 右侧：头像
        right_cell = table.cell(0, 1)
        right_cell.vertical_alignment = 1  # 居中
        if avatar_path and os.path.exists(avatar_path):
            p_avatar = right_cell.paragraphs[0]
            p_avatar.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_avatar = p_avatar.add_run()
            run_avatar.add_picture(avatar_path, width=Cm(3.2))

        # 去掉表格边框
        tbl = table._tbl
        tblPr = tbl.tblPr if tbl.tblPr is not None else tbl.makeelement(qn("w:tblPr"), {})
        borders = tblPr.makeelement(qn("w:tblBorders"), {})
        for border_name in ["top", "left", "bottom", "right", "insideH", "insideV"]:
            border = borders.makeelement(qn(f"w:{border_name}"), {
                qn("w:val"): "none",
                qn("w:sz"): "0",
                qn("w:space"): "0",
                qn("w:color"): "auto",
            })
            borders.append(border)
        tblPr.append(borders)

        # 头部底部分隔线
        p_line = doc.add_paragraph()
        p_line.paragraph_format.space_before = Pt(2)
        p_line.paragraph_format.space_after = Pt(6)
        pPr = p_line._element.get_or_add_pPr()
        pBdr = pPr.makeelement(qn("w:pBdr"), {})
        bottom = pBdr.makeelement(qn("w:bottom"), {
            qn("w:val"): "single",
            qn("w:sz"): "12",
            qn("w:space"): "1",
            qn("w:color"): "333333",
        })
        pBdr.append(bottom)
        pPr.append(pBdr)

    # ===== 继续解析剩下的 Markdown =====
    i = start_idx
    while i < len(lines):
        line = lines[i].rstrip("\n")

        # 空行
        if not line.strip():
            i += 1
            continue

        # 分隔线
        if line.strip() == "---":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run("─" * 50)
            set_font(run, size=8, color=(180, 180, 180))
            i += 1
            continue

        # H1（姓名）
        if line.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(line[2:].strip())
            set_font(run, size=18, bold=True)
            i += 1
            continue

        # H2（教育背景、专业技能等）
        if line.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(line[3:].strip())
            set_font(run, size=12, bold=True)
            # 底部边框
            pPr = p._element.get_or_add_pPr()
            pBdr = pPr.makeelement(qn("w:pBdr"), {})
            bottom = pBdr.makeelement(qn("w:bottom"), {
                qn("w:val"): "single",
                qn("w:sz"): "6",
                qn("w:space"): "1",
                qn("w:color"): "333333",
            })
            pBdr.append(bottom)
            pPr.append(pBdr)
            i += 1
            continue

        # H3（项目名称）
        if line.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(1)
            run = p.add_run(line[4:].strip())
            set_font(run, size=11, bold=True)
            i += 1
            continue

        # 无序列表
        if line.strip().startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.left_indent = Cm(0.5)
            add_inline_formatting(p, line.strip()[2:])
            i += 1
            continue

        # 普通段落（可能是 H3 下面的时间/技术栈行）
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        # 检测是否是居中的联系信息行（H1 后面的第一行）
        if i > 0 and lines[i-1].startswith("# "):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(line.strip())
            set_font(run, size=9, color=(100, 100, 100))
        else:
            add_inline_formatting(p, line.strip())
        i += 1

    # 保存
    try:
        os.makedirs(os.path.dirname(docx_path) or ".", exist_ok=True)
        doc.save(docx_path)
        print(f"✅ Word 已生成: {docx_path}")
        print(f"   大小: {os.path.getsize(docx_path) / 1024:.1f} KB")
        return True
    except Exception as e:
        print(f"❌ Word 生成失败: {e}", file=sys.stderr)
        return False


def main():
    parser = argparse.ArgumentParser(description="Markdown 简历 → Word（带头像）")
    parser.add_argument("input", help="输入 Markdown 文件路径")
    parser.add_argument("--output", "-o", help="输出 DOCX 路径（默认同目录同名）", default=None)
    parser.add_argument("--avatar", "-a", help="头像图片路径（默认使用 skill/assets/avatar.jpg）", default=None)
    parser.add_argument("--no-avatar", action="store_true", help="不包含头像")
    args = parser.parse_args()

    if args.output:
        docx_path = args.output
    else:
        docx_path = os.path.splitext(args.input)[0] + ".docx"

    avatar_path = None if args.no_avatar else (args.avatar or DEFAULT_AVATAR)
    success = md_to_docx(args.input, docx_path, avatar_path)
    sys.exit(0 if success else 1)


if __name__ == "__main__":
    main()
